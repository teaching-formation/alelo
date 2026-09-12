"""
documents.py — Génération de documents (PDF, Word, Excel) à partir d'un texte markdown.

alélo produit un contenu (réponse RAG) puis le rend dans le format demandé. Les fichiers
sont stockés temporairement et servis par l'API via /api/download/{id}.
"""

import io
import os
import re
import uuid
import tempfile
from datetime import datetime, timezone

DOC_DIR = os.getenv("DOC_DIR", os.path.join(tempfile.gettempdir(), "alelo_docs"))
os.makedirs(DOC_DIR, exist_ok=True)

FORMATS = {
    "pdf":  {"ext": "pdf",  "mime": "application/pdf", "label": "PDF"},
    "docx": {"ext": "docx", "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "label": "Word"},
    "xlsx": {"ext": "xlsx", "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "label": "Excel"},
}


# ── Parsing du contenu markdown ───────────────────────────────────────────────
def _slug(s: str) -> str:
    s = re.sub(r"[^\w\s-]", "", (s or "").strip().lower())
    s = re.sub(r"[\s_-]+", "-", s).strip("-")
    return s[:50] or "document"


def _strip_inline(s: str) -> str:
    """Retire le markdown inline pour du texte brut."""
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)     # [txt](url) → txt
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)             # gras
    s = re.sub(r"[*_`#>]", "", s)
    return s.strip()


def _bold_runs(s: str):
    """Découpe une ligne en segments (texte, gras) pour Word/PDF."""
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    parts = re.split(r"(\*\*.+?\*\*)", s)
    runs = []
    for p in parts:
        if not p:
            continue
        if p.startswith("**") and p.endswith("**"):
            runs.append((p[2:-2], True))
        else:
            runs.append((re.sub(r"[*_`]", "", p), False))
    return runs


def _parse(body: str):
    """Transforme le markdown en blocs (type, texte). type ∈ h/li_ul/li_ol/p."""
    blocks = []
    for raw in (body or "").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        h = re.match(r"^\s*(#{1,4})\s+(.*)", line)
        ul = re.match(r"^\s*[-*•]\s+(.*)", line)
        ol = re.match(r"^\s*\d+[.)]\s+(.*)", line)
        if h:
            blocks.append(("h", h.group(2).strip()))
        elif ul:
            blocks.append(("li_ul", ul.group(1).strip()))
        elif ol:
            blocks.append(("li_ol", ol.group(1).strip()))
        else:
            blocks.append(("p", line.strip()))
    return blocks


_FOOTER = "Document généré par alélo — L'IA publique au service du citoyen (Côte d'Ivoire)."


def _now_str() -> str:
    return datetime.now(timezone.utc).astimezone().strftime("%d/%m/%Y")


# ── Word (.docx) ──────────────────────────────────────────────────────────────
def build_docx(title: str, body: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    h = doc.add_heading(title, level=0)
    sub = doc.add_paragraph(f"alélo · {_now_str()}")
    sub.runs[0].italic = True
    sub.runs[0].font.color.rgb = RGBColor(0x6b, 0x72, 0x80)

    for kind, text in _parse(body):
        if kind == "h":
            doc.add_heading(_strip_inline(text), level=2)
        elif kind in ("li_ul", "li_ol"):
            p = doc.add_paragraph(style="List Bullet" if kind == "li_ul" else "List Number")
            for t, b in _bold_runs(text):
                r = p.add_run(t)
                r.bold = b
        else:
            p = doc.add_paragraph()
            for t, b in _bold_runs(text):
                r = p.add_run(t)
                r.bold = b

    doc.add_paragraph()
    foot = doc.add_paragraph(_FOOTER)
    foot.runs[0].italic = True
    foot.runs[0].font.size = Pt(8)
    foot.runs[0].font.color.rgb = RGBColor(0x8a, 0x8f, 0x98)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── PDF (.pdf) ────────────────────────────────────────────────────────────────
def build_pdf(title: str, body: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_title(_strip_inline(title)[:120])

    def _w(txt):
        # fpdf core fonts = latin-1 ; on remplace ce qui n'y passe pas, et on découpe les
        # mots trop longs (URLs…) sinon multi_cell lève « Not enough horizontal space ».
        txt = re.sub(r"\S{55,}", lambda m: " ".join(re.findall(r".{1,50}", m.group(0))), txt or "")
        return txt.encode("latin-1", "replace").decode("latin-1")

    def cell(txt, h=6):
        try:
            pdf.multi_cell(w=pdf.epw, h=h, text=_w(txt), new_x="LMARGIN", new_y="NEXT")
        except Exception:
            pass

    pdf.set_font("Helvetica", "B", 18)
    cell(title, 9)
    pdf.set_font("Helvetica", "I", 9)
    pdf.set_text_color(107, 114, 128)
    cell(f"alélo · {_now_str()}", 6)
    pdf.set_text_color(20, 20, 20)
    pdf.ln(3)

    for kind, text in _parse(body):
        if kind == "h":
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 13)
            cell(_strip_inline(text), 7)
        else:
            prefix = "  -  " if kind in ("li_ul", "li_ol") else ""
            pdf.set_font("Helvetica", "", 11)
            cell(prefix + _strip_inline(text), 6)

    pdf.ln(6)
    pdf.set_font("Helvetica", "I", 8)
    pdf.set_text_color(138, 143, 152)
    cell(_FOOTER, 5)

    return bytes(pdf.output())


# ── Excel (.xlsx) ─────────────────────────────────────────────────────────────
def _md_table(body: str):
    """Détecte un tableau markdown (| a | b |). Retourne (headers, rows) ou None."""
    lines = [l for l in body.split("\n") if l.strip().startswith("|")]
    if len(lines) < 2:
        return None
    def cells(l):
        return [c.strip() for c in l.strip().strip("|").split("|")]
    headers = cells(lines[0])
    rows = [cells(l) for l in lines[2:] if set(l.strip()) - set("|-: ")]  # saute la ligne ---
    return headers, rows


def build_xlsx(title: str, body: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "alélo"

    ws["A1"] = title
    ws["A1"].font = Font(size=15, bold=True)
    ws["A2"] = f"alélo · {_now_str()}"
    ws["A2"].font = Font(italic=True, color="6B7280")

    table = _md_table(body)
    row = 4
    if table:
        headers, rows = table
        for j, hdr in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=j, value=_strip_inline(hdr))
            cell.font = Font(bold=True)
        row += 1
        for r in rows:
            for j, val in enumerate(r, start=1):
                ws.cell(row=row, column=j, value=_strip_inline(val))
            row += 1
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 60)
    else:
        # prose → une ligne par bloc, titres en gras
        ws.column_dimensions["A"].width = 100
        for kind, text in _parse(body):
            cell = ws.cell(row=row, column=1, value=_strip_inline(text))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if kind == "h":
                cell.font = Font(bold=True, size=12)
            elif kind in ("li_ul", "li_ol"):
                cell.value = "• " + _strip_inline(text)
            row += 1

    ws.cell(row=row + 1, column=1, value=_FOOTER).font = Font(italic=True, size=8, color="8A8F98")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_BUILDERS = {"docx": build_docx, "pdf": build_pdf, "xlsx": build_xlsx}


# ── API publique ──────────────────────────────────────────────────────────────
def generate(fmt: str, title: str, body: str) -> dict:
    """Génère le document, l'écrit sur disque, renvoie {id, filename, format, label, mime}."""
    fmt = (fmt or "pdf").lower()
    if fmt not in FORMATS:
        fmt = "pdf"
    data = _BUILDERS[fmt](title or "Document", body or "")
    file_id = uuid.uuid4().hex
    ext = FORMATS[fmt]["ext"]
    filename = f"alelo-{_slug(title)}.{ext}"
    # Le nom d'affichage est encodé dans le nom sur disque ({id}__{filename})
    # car pour un téléchargement cross-origin, seul le Content-Disposition serveur fait foi.
    with open(os.path.join(DOC_DIR, f"{file_id}__{filename}"), "wb") as f:
        f.write(data)
    return {"id": file_id, "filename": filename, "format": fmt,
            "label": FORMATS[fmt]["label"], "mime": FORMATS[fmt]["mime"]}


def get_file(file_id: str):
    """Retourne (path, mime, filename) pour un id, ou None."""
    if not re.fullmatch(r"[0-9a-f]{32}", file_id or ""):
        return None
    for name in os.listdir(DOC_DIR):
        if name.startswith(f"{file_id}__"):
            path = os.path.join(DOC_DIR, name)
            filename = name.split("__", 1)[1]
            ext = filename.rsplit(".", 1)[-1].lower()
            mime = next((m["mime"] for m in FORMATS.values() if m["ext"] == ext),
                        "application/octet-stream")
            return path, mime, filename
    return None
